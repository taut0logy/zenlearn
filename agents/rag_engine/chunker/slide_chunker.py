"""
Slide (PPTX) Chunker - Content-aware chunking for PowerPoint presentations.
"""

from dataclasses import dataclass
from typing import List, Dict, Any, Optional
import os

from .base import BaseChunker, Chunk, ChunkingConfig, count_tokens

# Try to import python-pptx
try:
    from pptx import Presentation
    from pptx.util import Inches
    HAS_PPTX = True
except ImportError:
    HAS_PPTX = False


@dataclass
class SlideContent:
    """Extracted content from a slide."""
    slide_number: int
    title: Optional[str]
    body_text: str
    speaker_notes: Optional[str]
    has_code: bool = False
    has_images: bool = False


class SlideChunker(BaseChunker):
    """
    Chunker for PowerPoint presentations.
    
    Strategy:
    - Each slide is a base unit
    - Merge small adjacent slides
    - Preserve code slides intact
    - Include speaker notes
    """
    
    def __init__(self, config: Optional[ChunkingConfig] = None):
        default_config = ChunkingConfig(
            min_tokens=100,
            max_tokens=800,
            target_tokens=400,
            overlap_tokens=50
        )
        super().__init__(config or default_config)
    
    def chunk(
        self, 
        content: Any, 
        metadata: Dict[str, Any]
    ) -> List[Chunk]:
        """
        Chunk a PPTX file.
        
        Args:
            content: Path to PPTX file
            metadata: Base metadata (source_file, course_id, etc.)
            
        Returns:
            List of chunks
        """
        if not HAS_PPTX:
            raise ImportError("python-pptx required: pip install python-pptx")
        
        file_path = content
        slides = self._extract_slides(file_path)
        
        return self._chunk_slides(slides, metadata)
    
    def chunk_from_slides(
        self,
        slides: List[SlideContent],
        metadata: Dict[str, Any]
    ) -> List[Chunk]:
        """Chunk from already-extracted slides."""
        return self._chunk_slides(slides, metadata)
    
    def _extract_slides(self, file_path: str) -> List[SlideContent]:
        """Extract content from PPTX file."""
        prs = Presentation(file_path)
        slides = []
        
        for idx, slide in enumerate(prs.slides, 1):
            title = None
            body_parts = []
            speaker_notes = None
            has_code = False
            has_images = False
            
            for shape in slide.shapes:
                # Extract title
                if shape.is_placeholder:
                    if hasattr(shape, 'placeholder_format'):
                        if shape.placeholder_format.type == 1:  # Title
                            if shape.has_text_frame:
                                title = shape.text_frame.text.strip()
                
                # Extract text
                if shape.has_text_frame:
                    text = self._extract_text_frame(shape.text_frame)
                    if text and text != title:
                        body_parts.append(text)
                        if self._looks_like_code(text):
                            has_code = True
                
                # Check for images
                if hasattr(shape, 'image'):
                    has_images = True
            
            # Speaker notes
            if slide.has_notes_slide:
                notes_frame = slide.notes_slide.notes_text_frame
                if notes_frame:
                    speaker_notes = notes_frame.text.strip()
            
            slides.append(SlideContent(
                slide_number=idx,
                title=title,
                body_text="\n\n".join(body_parts),
                speaker_notes=speaker_notes,
                has_code=has_code,
                has_images=has_images
            ))
        
        return slides
    
    def _chunk_slides(
        self, 
        slides: List[SlideContent],
        base_metadata: Dict[str, Any]
    ) -> List[Chunk]:
        """Apply chunking strategy to slides."""
        chunks = []
        buffer = []
        buffer_tokens = 0
        current_section = None
        
        for slide in slides:
            slide_text = self._format_slide(slide)
            slide_tokens = count_tokens(slide_text)
            
            # Code slides stay intact
            if slide.has_code:
                # Flush buffer first
                if buffer:
                    chunks.append(self._create_chunk(buffer, current_section, base_metadata))
                    buffer = []
                    buffer_tokens = 0
                
                # Add code slide as own chunk
                chunks.append(self._create_chunk([slide], current_section, base_metadata))
                continue
            
            # Check if adding this slide exceeds max
            if buffer_tokens + slide_tokens > self.config.max_tokens:
                if buffer:
                    chunks.append(self._create_chunk(buffer, current_section, base_metadata))
                buffer = [slide]
                buffer_tokens = slide_tokens
            else:
                buffer.append(slide)
                buffer_tokens += slide_tokens
            
            # Update section from titles
            if slide.title and self._is_section_header(slide):
                current_section = slide.title
        
        # Flush remaining buffer
        if buffer:
            chunks.append(self._create_chunk(buffer, current_section, base_metadata))
        
        return chunks
    
    def _create_chunk(
        self,
        slides: List[SlideContent],
        section: Optional[str],
        base_metadata: Dict[str, Any]
    ) -> Chunk:
        """Create a chunk from slides."""
        content = "\n\n---\n\n".join(self._format_slide(s) for s in slides)
        
        # Slide range for citation
        if len(slides) == 1:
            slide_ref = f"Slide {slides[0].slide_number}"
        else:
            slide_ref = f"Slides {slides[0].slide_number}-{slides[-1].slide_number}"
        
        chunk_metadata = self._merge_metadata(base_metadata, {
            "chunk_type": "slides",
            "slide_numbers": [s.slide_number for s in slides],
            "slide_range": slide_ref,
            "section": section,
            "titles": [s.title for s in slides if s.title],
            "has_code": any(s.has_code for s in slides),
        })
        
        return Chunk(
            content=content,
            metadata=chunk_metadata,
            token_count=count_tokens(content)
        )
    
    def _format_slide(self, slide: SlideContent) -> str:
        """Format slide content for embedding."""
        parts = []
        
        if slide.title:
            parts.append(f"## Slide {slide.slide_number}: {slide.title}")
        else:
            parts.append(f"## Slide {slide.slide_number}")
        
        if slide.body_text:
            parts.append(slide.body_text)
        
        if slide.speaker_notes:
            parts.append(f"\n[Speaker Notes:]\n{slide.speaker_notes}")
        
        return "\n\n".join(parts)
    
    def _extract_text_frame(self, text_frame) -> str:
        """Extract text preserving structure."""
        paragraphs = []
        for para in text_frame.paragraphs:
            text = para.text.strip()
            if text:
                if para.level > 0:
                    indent = "  " * para.level
                    text = f"{indent}• {text}"
                paragraphs.append(text)
        return "\n".join(paragraphs)
    
    def _looks_like_code(self, text: str) -> bool:
        """Check if text appears to be code."""
        indicators = ['def ', 'class ', 'import ', 'function ', 'const ', '= {', '();']
        return any(ind in text for ind in indicators)
    
    def _is_section_header(self, slide: SlideContent) -> bool:
        """Check if slide is a section header."""
        if not slide.title:
            return False
        body_words = len(slide.body_text.split()) if slide.body_text else 0
        return body_words < 20


# Factory function
def get_slide_chunker(config: Optional[ChunkingConfig] = None) -> SlideChunker:
    return SlideChunker(config)
